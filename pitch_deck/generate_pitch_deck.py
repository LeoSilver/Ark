"""ネガティブ・チャーン & LTV向上の説得ロジック — ピッチ資料生成スクリプト

python-pptx で 16:9 / 全7スライドの .pptx を生成する。
    pip install python-pptx
    python generate_pitch_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "pitch_deck_negative_churn.pptx"
LOGO_ON_DARK = BASE_DIR / "assets" / "taxel_logo_dark.png"
LOGO_ON_LIGHT = BASE_DIR / "assets" / "taxel_logo_light.png"

# ---- カラーパレット ----
NAVY = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x25, 0x63, 0xEB)
CARD = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PALE_BLUE = RGBColor(0xDB, 0xEA, 0xFE)
NAVY_CARD = RGBColor(0x1E, 0x29, 0x3B)
DARK_MUTED = RGBColor(0x94, 0xA3, 0xB8)

FONT = "Meiryo"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6


# ---------------------------------------------------------------- helpers
def set_font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=16, color=NAVY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """lines: str または [(text, size, color, bold), ...] / 文字列のリスト"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if isinstance(line, tuple):
            t, s, c, b = line
        else:
            t, s, c, b = line, size, color, bold
        run = p.add_run()
        run.text = t
        set_font(run, s, c, b)
    return box


def rect(slide, x, y, w, h, fill, line=None, rounded=True, radius=0.08):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def badge(slide, x, y, d, label, fill=BLUE, color=WHITE, size=14):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    circle.fill.solid()
    circle.fill.fore_color.rgb = fill
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    set_font(run, size, color, True)
    return circle


def arrow(slide, x, y, w, h, fill=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def logo(slide, path, x, y, h):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))


def page_number(slide, n, color=MUTED):
    text(slide, SLIDE_W - MARGIN - 0.5, SLIDE_H - 0.45, 0.5, 0.25, str(n),
         size=10, color=color, align=PP_ALIGN.RIGHT)


def content_header(slide, n, kicker, title, size=28):
    """ライトスライド共通ヘッダー：キッカー＋タイトル＋右上ロゴ"""
    background(slide, WHITE)
    text(slide, MARGIN, 0.45, 9.5, 0.3, kicker, size=12, color=BLUE, bold=True)
    text(slide, MARGIN, 0.78, 10.6, 0.9, title, size=size, color=NAVY, bold=True)
    logo(slide, LOGO_ON_LIGHT, SLIDE_W - MARGIN - 1.45, 0.45, 0.4)
    page_number(slide, n)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # Blank


# ---------------------------------------------------------------- slides
def slide1_title(prs):
    s = new_slide(prs)
    background(s, NAVY)
    logo(s, LOGO_ON_DARK, MARGIN + 0.1, 0.6, 0.62)
    text(s, MARGIN + 0.1, 2.05, 11, 0.35, "INVESTOR BRIEFING ｜ UNIT ECONOMICS",
         size=13, color=BLUE, bold=True)
    text(s, MARGIN + 0.1, 2.5, 12.1, 1.1, "ネガティブ・チャーン & LTV向上の説得ロジック",
         size=34, color=WHITE, bold=True)
    text(s, MARGIN + 0.1, 3.6, 11, 0.5, "税理士チャネルが生む「構造的な」収益拡大の4つの根拠",
         size=18, color=DARK_MUTED)

    axes = ["構造的クロスセル", "スイッチングコスト", "LTV / CAC 拡大", "シナリオ分析"]
    w, gap, y = 2.85, 0.2, 5.15
    for i, label in enumerate(axes):
        x = MARGIN + 0.1 + i * (w + gap)
        rect(s, x, y, w, 0.9, NAVY_CARD)
        badge(s, x + 0.2, y + 0.2, 0.5, str(i + 1), size=14)
        text(s, x + 0.85, y, w - 0.95, 0.9, label, size=14, color=WHITE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    text(s, MARGIN + 0.1, 6.65, 8, 0.3, "taXel（タクセル）｜ 投資家向け補足資料 ｜ 2026年9月",
         size=11, color=DARK_MUTED)


def slide2_overview(prs):
    s = new_slide(prs)
    content_header(s, 2, "OVERVIEW", "投資家が求める「Why & How」の証明")

    # 投資家の問い
    rect(s, MARGIN, 1.85, SLIDE_W - 2 * MARGIN, 0.85, PALE_BLUE)
    text(s, MARGIN + 0.35, 1.85, 1.6, 0.85, "投資家の問い", size=14, color=BLUE, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, MARGIN + 2.0, 1.85, 9.8, 0.85,
         "なぜ解約が起きても、既存顧客からの売上は伸び続けるのか？",
         size=18, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    cards = [
        ("HOW", "構造的クロスセル", "税理士経由で利用プロダクト数が拡大"),
        ("WHY", "スイッチングコスト", "業務インフラ化＋税理士CSで高定着"),
        ("PROOF", "LTV / CAC 拡大", "初年度3倍 → 3年目6倍以上"),
        ("RISK", "シナリオ分析", "保守的な下限とアップサイドを両提示"),
    ]
    w, gap, y, h = 2.85, 0.2, 3.05, 2.55
    for i, (tag, title, body) in enumerate(cards):
        x = MARGIN + i * (w + gap)
        rect(s, x, y, w, h, CARD, line=BORDER)
        badge(s, x + 0.3, y + 0.3, 0.6, f"{i + 1}", size=16)
        text(s, x + 1.05, y + 0.3, 1.6, 0.6, f"軸{i + 1} ｜ {tag}", size=12, color=BLUE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, y + 1.1, w - 0.6, 0.45, title, size=16, bold=True)
        text(s, x + 0.3, y + 1.6, w - 0.6, 0.8, body, size=13, color=MUTED)

    text(s, MARGIN, 6.0, SLIDE_W - 2 * MARGIN, 0.5,
         [("ネガティブ・チャーンは「期待値」ではなく「構造」から生まれる", 18, NAVY, True)],
         align=PP_ALIGN.CENTER)


def slide3_crosssell(prs):
    s = new_slide(prs)
    content_header(s, 3, "軸1 ｜ HOW",
                   "構造的なクロスセルの仕組み（税理士チャネルの横展開）")

    # フロー：4ステップ
    steps = [
        ("税理士事務所を獲得", "チャネルの起点"),
        ("1プロダクトで導入", "小さく始める"),
        ("信頼関係の構築", "運用実績が蓄積"),
        ("他プロダクトへ展開", "同一税理士経由"),
    ]
    gap, y, h = 0.52, 1.95, 1.45
    w = (SLIDE_W - 2 * MARGIN - 3 * gap) / 4
    for i, (title, sub) in enumerate(steps):
        x = MARGIN + i * (w + gap)
        last = i == len(steps) - 1
        rect(s, x, y, w, h, BLUE if last else CARD, line=None if last else BORDER)
        text(s, x + 0.25, y + 0.22, w - 0.5, 0.3, f"STEP {i + 1}", size=11,
             color=WHITE if last else BLUE, bold=True)
        text(s, x + 0.25, y + 0.55, w - 0.5, 0.45, title, size=16,
             color=WHITE if last else NAVY, bold=True)
        text(s, x + 0.25, y + 0.98, w - 0.5, 0.3, sub, size=12,
             color=PALE_BLUE if last else MUTED)
        if not last:
            arrow(s, x + w + 0.1, y + h / 2 - 0.16, 0.32, 0.32)

    # 下段左：1税理士あたりのプロダクト拡大イメージ（積み上げブロック）
    px, py, pw, ph = MARGIN, 3.8, 6.4, 2.95
    rect(s, px, py, pw, ph, CARD, line=BORDER)
    text(s, px + 0.3, py + 0.2, pw - 0.6, 0.3, "1税理士あたり利用プロダクト数（イメージ）",
         size=13, bold=True)
    years = ["1年目", "2年目", "3年目"]
    products = ["プロダクトA", "＋プロダクトB", "＋プロダクトC"]
    shades = [BLUE, RGBColor(0x60, 0x8F, 0xF2), RGBColor(0x93, 0xB4, 0xF7)]
    col_w, base_y, blk_h = 1.45, py + ph - 0.55, 0.52
    for i, yr in enumerate(years):
        cx = px + 0.55 + i * (col_w + 0.45)
        for j in range(i + 1):
            b = rect(s, cx, base_y - (j + 1) * (blk_h + 0.06), col_w, blk_h, shades[j], radius=0.12)
            tf = b.text_frame
            tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = products[j].lstrip("＋") if j == 0 else products[j]
            set_font(r, 11, WHITE, True)
        text(s, cx, base_y + 0.1, col_w, 0.3, yr, size=12, color=MUTED, align=PP_ALIGN.CENTER)

    # 下段右：結論
    rx, rw = px + pw + 0.35, SLIDE_W - MARGIN - (px + pw + 0.35)
    rect(s, rx, py, rw, 1.3, CARD, line=BORDER)
    text(s, rx + 0.3, py + 0.2, rw - 0.6, 0.3, "単一プロダクトの解約", size=13, color=MUTED, bold=True)
    text(s, rx + 0.3, py + 0.55, rw - 0.6, 0.6, "発生しても影響は限定的", size=18, bold=True)
    rect(s, rx, py + 1.5, rw, 1.45, NAVY)
    text(s, rx + 0.3, py + 1.68, rw - 0.6, 0.3, "1税理士あたり ARPU・利用プロダクト数",
         size=13, color=DARK_MUTED, bold=True)
    text(s, rx + 0.3, py + 2.02, rw - 0.6, 0.4, "年々拡大", size=20, color=WHITE, bold=True)
    text(s, rx + 0.3, py + 2.45, rw - 0.6, 0.45, "＝ 全体でネガティブ・チャーンを達成",
         size=18, color=RGBColor(0x93, 0xB4, 0xF7), bold=True)


def slide4_switching(prs):
    s = new_slide(prs)
    content_header(s, 4, "軸2 ｜ WHY",
                   "スイッチングコストと高定着性の根拠")

    colw, gap, y, h = 5.9, 0.33, 1.95, 4.75
    lx, rx = MARGIN, MARGIN + colw + gap

    # 左：業務インフラ化
    rect(s, lx, y, colw, h, CARD, line=BORDER)
    badge(s, lx + 0.35, y + 0.35, 0.6, "A", size=16)
    text(s, lx + 1.15, y + 0.35, colw - 1.5, 0.6, "業務インフラ化", size=20, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, lx + 0.35, y + 1.15, colw - 0.7, 0.35, "導入6か月超で、日常業務フローに完全定着",
         size=14, color=MUTED)
    # タイムライン
    ty = y + 2.05
    line = s.shapes.add_connector(1, Inches(lx + 0.6), Inches(ty + 0.35),
                                  Inches(lx + colw - 0.6), Inches(ty + 0.35))
    line.line.color.rgb = BORDER
    line.line.width = Pt(3)
    for label, frac, active in [("導入", 0.0, False), ("6か月", 0.5, True), ("定着", 1.0, False)]:
        cx = lx + 0.6 + frac * (colw - 1.2)
        badge(s, cx - 0.15, ty + 0.2, 0.3, "", fill=BLUE if active else MUTED)
        text(s, cx - 0.7, ty - 0.2, 1.4, 0.3, label, size=12, color=BLUE if active else MUTED,
             bold=True, align=PP_ALIGN.CENTER)
    chips = ["経理", "決済", "リスク管理"]
    cw = (colw - 0.7 - 0.2 * 2) / 3
    for i, c in enumerate(chips):
        b = rect(s, lx + 0.35 + i * (cw + 0.2), y + 3.0, cw, 0.5, PALE_BLUE, radius=0.5)
        tf = b.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = c
        set_font(r, 13, BLUE, True)
    text(s, lx + 0.35, y + 3.75, colw - 0.7, 0.7,
         [("スイッチングコスト上昇 → 解約率が急減", 16, NAVY, True)])

    # 右：税理士によるサクセスサポート
    rect(s, rx, y, colw, h, CARD, line=BORDER)
    badge(s, rx + 0.35, y + 0.35, 0.6, "B", size=16)
    text(s, rx + 1.15, y + 0.35, colw - 1.5, 0.6, "税理士によるサクセスサポート", size=20,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, rx + 0.35, y + 1.15, colw - 0.7, 0.35, "顧問税理士が「実質的なCS担当」として機能",
         size=14, color=MUTED)
    points = [
        ("決算・財務を把握", "顧客の状況を最も理解する立場"),
        ("継続的に関与・提案", "定例の顧問業務の中で活用を促進"),
        ("CSコスト抑制 × 高定着", "自社CSに依存しない定着モデル"),
    ]
    for i, (t, b) in enumerate(points):
        py = y + 1.75 + i * 0.95
        rect(s, rx + 0.35, py, colw - 0.7, 0.8, WHITE, line=BORDER)
        text(s, rx + 0.6, py + 0.1, colw - 1.2, 0.32, t, size=14, bold=True)
        text(s, rx + 0.6, py + 0.44, colw - 1.2, 0.3, b, size=12, color=MUTED)


def slide5_unit_economics(prs):
    s = new_slide(prs)
    content_header(s, 5, "軸3 ｜ PROOF",
                   "ユニットエコノミクス（LTV / CAC）の拡大ストーリー")

    # 左：チャート
    cx, cy, cw, ch = MARGIN, 1.95, 5.4, 4.75
    rect(s, cx, cy, cw, ch, CARD, line=BORDER)
    text(s, cx + 0.3, cy + 0.25, cw - 0.6, 0.3, "LTV / CAC 倍率", size=14, bold=True)
    text(s, cx + 0.3, cy + 0.6, cw - 0.6, 0.3, "初年度 → 3年目（6倍以上を目標）", size=12, color=MUTED)
    data = CategoryChartData()
    data.categories = ["初年度", "3年目"]
    data.add_series("LTV/CAC", (3, 6))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(cx + 0.3), Inches(cy + 1.0),
                            Inches(cw - 0.6), Inches(ch - 1.25), data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 90
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0"倍"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(20)
    dl.font.bold = True
    dl.font.color.rgb = NAVY
    dl.font.name = FONT
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = BLUE
    pt = series.points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = RGBColor(0x93, 0xB4, 0xF7)
    va = chart.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 7.5
    va.minimum_scale = 0
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(13)
    ca.tick_labels.font.color.rgb = MUTED
    ca.tick_labels.font.name = FONT
    ca.format.line.color.rgb = BORDER
    ca.has_major_gridlines = False
    # 3年目ラベルに「以上」を補足
    text(s, cx + cw - 2.45, cy + 1.05, 2.0, 0.3, "6倍以上へ", size=12, color=BLUE, bold=True,
         align=PP_ALIGN.CENTER)

    # 右：ドライバー
    rx = cx + cw + 0.4
    rw = SLIDE_W - MARGIN - rx
    text(s, rx, 1.95, rw, 0.35, "倍率を押し上げる3つのドライバー", size=14, color=BLUE, bold=True)
    drivers = [
        ("CAC", "税理士チャネル活用で劇的に低位"),
        ("初期解約率", "導入初期の離脱を抑制"),
        ("クロスセル", "1社あたりLTVを継続的に拡大"),
    ]
    dw = (rw - 0.4) / 3
    for i, (t, b) in enumerate(drivers):
        x = rx + i * (dw + 0.2)
        rect(s, x, 2.4, dw, 1.85, CARD, line=BORDER)
        text(s, x + 0.25, 2.6, dw - 0.5, 0.4, t, size=17, color=BLUE, bold=True)
        text(s, x + 0.25, 3.1, dw - 0.5, 1.0, b, size=13)

    # Payback
    rect(s, rx, 4.5, rw, 2.2, NAVY)
    text(s, rx + 0.35, 4.7, rw - 0.7, 0.3, "PAYBACK PERIOD", size=12, color=DARK_MUTED, bold=True)
    text(s, rx + 0.35, 5.05, rw - 0.7, 0.5, "CAC回収期間が極めて短い", size=20, color=WHITE, bold=True)
    text(s, rx + 0.35, 5.7, rw - 0.7, 0.8, "回収後の利益がストックとして積み上がる",
         size=15, color=RGBColor(0x93, 0xB4, 0xF7), bold=True)


def slide6_scenarios(prs):
    s = new_slide(prs)
    content_header(s, 6, "軸4 ｜ RISK",
                   "シナリオ分析による実現可能性（ベースライン vs アップサイド）", size=22)

    colw, gap, y, h = 3.6, 0.3, 1.95, 4.75
    # 左2列：シナリオカード
    specs = [
        ("ベースライン", "BASE CASE", CARD, NAVY, MUTED,
         ["解約率 2.0% 固定", "クロスセル なし", "保守的な下限値"]),
        ("アップサイド", "UPSIDE CASE", BLUE, WHITE, PALE_BLUE,
         ["初期解約率の低下", "税理士経由のクロスセル", "実効LTV ＋30〜50%"]),
    ]
    for i, (title, sub, fill, fg, sub_c, items) in enumerate(specs):
        x = MARGIN + i * (colw + gap)
        rect(s, x, y, colw, h, fill, line=BORDER if fill == CARD else None)
        text(s, x + 0.35, y + 0.35, colw - 0.7, 0.3, sub, size=12, color=sub_c, bold=True)
        text(s, x + 0.35, y + 0.7, colw - 0.7, 0.5, title, size=22, color=fg, bold=True)
        for j, item in enumerate(items):
            iy = y + 1.55 + j * 0.95
            badge(s, x + 0.35, iy + 0.08, 0.36, "✓", fill=BLUE if fill == CARD else WHITE,
                  color=WHITE if fill == CARD else BLUE, size=11)
            text(s, x + 0.9, iy, colw - 1.2, 0.55, item, size=15, color=fg, bold=j == 2,
                 anchor=MSO_ANCHOR.MIDDLE)

    # 右：実効LTV比較（シェイプで描く指数バー）
    vx = MARGIN + 2 * (colw + gap)
    vw = SLIDE_W - MARGIN - vx
    rect(s, vx, y, vw, h, CARD, line=BORDER)
    text(s, vx + 0.35, y + 0.35, vw - 0.7, 0.3, "実効LTV（ベースライン＝100）", size=14, bold=True)
    base_y = y + h - 0.75
    unit = 2.45 / 150  # 150 → 2.45in
    bw = 1.2
    bx1, bx2 = vx + 0.9, vx + vw - 0.9 - bw
    # ベースライン 100
    rect(s, bx1, base_y - 100 * unit, bw, 100 * unit, RGBColor(0x94, 0xA3, 0xB8), rounded=False)
    text(s, bx1 - 0.2, base_y - 100 * unit - 0.45, bw + 0.4, 0.4, "100", size=18, bold=True,
         align=PP_ALIGN.CENTER)
    # アップサイド 130（確度高）+ 20（上限まで）
    rect(s, bx2, base_y - 130 * unit, bw, 130 * unit, BLUE, rounded=False)
    rect(s, bx2, base_y - 150 * unit, bw, 20 * unit, PALE_BLUE, line=BLUE, rounded=False)
    text(s, bx2 - 0.3, base_y - 150 * unit - 0.45, bw + 0.6, 0.4, "130〜150", size=18,
         color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    for bx, lab in [(bx1, "ベースライン"), (bx2, "アップサイド")]:
        text(s, bx - 0.3, base_y + 0.12, bw + 0.6, 0.3, lab, size=12, color=MUTED,
             align=PP_ALIGN.CENTER)
    text(s, vx + 0.35, y + 0.75, vw - 0.7, 0.6, "ベースラインを下限に、上振れ余地を提示",
         size=12, color=MUTED)


def slide7_conclusion(prs):
    s = new_slide(prs)
    background(s, NAVY)
    logo(s, LOGO_ON_DARK, SLIDE_W - MARGIN - 1.6, 0.5, 0.45)
    page_number(s, 7, DARK_MUTED)
    text(s, MARGIN + 0.1, 0.55, 8, 0.3, "CONCLUSION", size=12, color=BLUE, bold=True)
    text(s, MARGIN + 0.1, 0.95, 12, 1.4,
         [("単体SaaSではなく", 22, DARK_MUTED, True),
          ("「税理士プラットフォーム」としての構造的優位性", 34, WHITE, True)],
         spacing=1.2)

    cards = [
        ("獲得", "低CAC", "税理士チャネルで効率的に獲得"),
        ("定着", "高い定着性", "業務インフラ化＋税理士CS"),
        ("拡張", "クロスセル", "同一税理士経由で横展開"),
        ("結果", "ネガティブ・チャーン", "LTVが時間とともに拡大"),
    ]
    gap, y, h = 0.55, 2.95, 2.3
    x0 = MARGIN + 0.1
    total = SLIDE_W - MARGIN - x0
    widths = [2.35, 2.35, 2.35]
    widths.append(total - sum(widths) - 3 * gap)
    x = x0
    for i, ((kw, title, body), w) in enumerate(zip(cards, widths)):
        last = i == len(cards) - 1
        rect(s, x, y, w, h, BLUE if last else NAVY_CARD)
        text(s, x + 0.3, y + 0.3, w - 0.6, 0.35, kw, size=14,
             color=PALE_BLUE if last else BLUE, bold=True)
        text(s, x + 0.3, y + 0.75, w - 0.6, 0.5, title, size=18, color=WHITE, bold=True)
        text(s, x + 0.3, y + 1.4, w - 0.6, 0.7, body, size=13,
             color=PALE_BLUE if last else DARK_MUTED)
        if not last:
            arrow(s, x + w + 0.12, y + h / 2 - 0.16, 0.3, 0.32)
        x += w + gap

    rect(s, MARGIN + 0.1, 5.75, SLIDE_W - 2 * MARGIN - 0.1, 0.95, NAVY_CARD)
    text(s, MARGIN + 0.45, 5.75, SLIDE_W - 2 * MARGIN - 0.8, 0.95,
         "税理士1事務所の獲得が、顧問先×複数プロダクトの収益へ波及",
         size=18, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for build in (slide1_title, slide2_overview, slide3_crosssell, slide4_switching,
                  slide5_unit_economics, slide6_scenarios, slide7_conclusion):
        build(prs)
    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
